from dotenv import load_dotenv
from pathlib import Path
ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

import os
import io
import re
import uuid
import logging
from datetime import datetime, timezone, timedelta
from typing import List, Optional, Dict, Any

import bcrypt
import jwt
import hmac
import hashlib
import base64
from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends, UploadFile, File, Query
from fastapi.responses import StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, EmailStr, Field
import openpyxl

# ----------------- Setup -----------------
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI(title="AP GROUP KPI Management")
api = APIRouter(prefix="/api")

JWT_ALGO = "HS256"

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("kpi")

# ----------------- Utils -----------------
def hash_pw(pw: str) -> str:
    return bcrypt.hashpw(pw.encode(), bcrypt.gensalt()).decode()

def verify_pw(pw: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(pw.encode(), hashed.encode())
    except Exception:
        return False

def make_token(sub: str, email: str, role: str, minutes: int = 60*24*7) -> str:
    payload = {
        "sub": sub, "email": email, "role": role,
        "exp": datetime.now(timezone.utc) + timedelta(minutes=minutes),
        "type": "access",
    }
    return jwt.encode(payload, os.environ["JWT_SECRET"], algorithm=JWT_ALGO)

async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        h = request.headers.get("Authorization", "")
        if h.startswith("Bearer "):
            token = h[7:]
    if not token:
        raise HTTPException(401, "Not authenticated")
    try:
        payload = jwt.decode(token, os.environ["JWT_SECRET"], algorithms=[JWT_ALGO])
    except jwt.ExpiredSignatureError:
        raise HTTPException(401, "Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(401, "Invalid token")
    user = await db.users.find_one({"id": payload["sub"]}, {"password_hash": 0})
    if not user:
        raise HTTPException(401, "User not found")
    user.pop("_id", None)
    return user

def require_role(*roles):
    async def dep(user: dict = Depends(get_current_user)):
        if user.get("role") not in roles:
            raise HTTPException(403, f"Access denied. Requires role: {roles}")
        return user
    return dep

def new_id() -> str:
    return str(uuid.uuid4())

def strip_id(d):
    if not d: return d
    d.pop("_id", None)
    return d

# ----------------- Models -----------------
class LoginReq(BaseModel):
    email: EmailStr
    password: str

class UserCreate(BaseModel):
    email: EmailStr
    password: str
    name: str
    role: str  # admin | supervisor | karyawan
    nik: Optional[str] = None
    divisi: Optional[str] = None  # required for supervisor role

class Divisi(BaseModel):
    id: Optional[str] = None
    kode: str
    nama: str

class Jabatan(BaseModel):
    id: Optional[str] = None
    kode: str
    nama: str
    divisi: str

class Karyawan(BaseModel):
    id: Optional[str] = None
    nik: str
    nama: str
    divisi: str
    jabatan: str
    atasan: Optional[str] = "Owner"
    status: str = "Aktif"
    # payroll master fields
    gaji_pokok: Optional[float] = 0
    tunjangan_transportasi: Optional[float] = 0
    tunjangan_makan: Optional[float] = 0
    tunjangan_kesehatan: Optional[float] = 0
    sistem_kerja: Optional[str] = "Bulanan"
    no_rek: Optional[str] = ""
    bank: Optional[str] = "BCA"
    no_hp: Optional[str] = ""

class KPIMaster(BaseModel):
    id: Optional[str] = None
    divisi: str
    kode: str
    nama: str
    satuan: str
    bobot: float
    arah_nilai: str  # Higher | Lower
    jenis_target: str  # Rasio | Kumulatif

class KPITarget(BaseModel):
    id: Optional[str] = None
    kode_kpi: str
    tahun: int
    bulan: int  # 1..12
    target: float

class KPIInput(BaseModel):
    id: Optional[str] = None
    bulan: int
    tahun: int
    cut_off: Optional[str] = None
    nik: str
    kode_kpi: str
    realisasi: float
    status: Optional[str] = None  # draft | submitted | approved | rejected; None = auto by role

class PayrollInput(BaseModel):
    id: Optional[str] = None
    bulan: int
    tahun: int
    nik: str
    periode_gaji: Optional[str] = ""
    lembur_jam: Optional[float] = 0
    terlambat_jam: Optional[float] = 0
    tidak_masuk_hari: Optional[float] = 0
    potongan_pinjaman: Optional[float] = 0
    potongan_lainnya: Optional[float] = 0
    bonus_lainnya: Optional[float] = 0
    catatan: Optional[str] = ""

class SettingModel(BaseModel):
    threshold_a: float = 95
    threshold_b: float = 85
    threshold_c: float = 75
    threshold_d: float = 65
    on_track_min: float = 85
    # payroll rates (per satuan)
    rate_lembur: float = 12000
    rate_terlambat: float = 12000
    rate_tidak_masuk: float = 36000

MONTHS_ID = ["Januari","Februari","Maret","April","Mei","Juni","Juli","Agustus","September","Oktober","November","Desember"]
MONTHS_SHORT = ["Jan","Feb","Mar","Apr","Mei","Jun","Jul","Agu","Sep","Okt","Nov","Des"]

# ----------------- Auth Endpoints -----------------
@api.post("/auth/login")
async def login(req: LoginReq, response: Response):
    email = req.email.lower()
    user = await db.users.find_one({"email": email})
    if not user or not verify_pw(req.password, user["password_hash"]):
        raise HTTPException(401, "Email atau password salah")
    token = make_token(user["id"], user["email"], user["role"])
    response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none", max_age=60*60*24*7, path="/")
    return {"id": user["id"], "email": user["email"], "name": user["name"], "role": user["role"],
            "nik": user.get("nik"), "divisi": user.get("divisi"), "token": token}

@api.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}

@api.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user

@api.post("/auth/register")
async def register_user(payload: UserCreate, user: dict = Depends(require_role("admin"))):
    email = payload.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(400, "Email sudah terdaftar")
    if payload.role == "supervisor" and not payload.divisi:
        raise HTTPException(400, "Supervisor wajib memilih divisi yang dinaungi")
    if payload.role == "karyawan" and not payload.nik:
        raise HTTPException(400, "Karyawan wajib dihubungkan dengan NIK")
    if payload.role == "supervisor":
        exists = await db.divisi.find_one({"nama": payload.divisi})
        if not exists:
            raise HTTPException(400, f"Divisi '{payload.divisi}' tidak ditemukan")
    if payload.role == "karyawan":
        exists = await db.karyawan.find_one({"nik": payload.nik})
        if not exists:
            raise HTTPException(400, f"NIK '{payload.nik}' tidak ditemukan di Master Karyawan")
    doc = {"id": new_id(), "email": email, "password_hash": hash_pw(payload.password),
           "name": payload.name, "role": payload.role, "nik": payload.nik,
           "divisi": payload.divisi if payload.role == "supervisor" else None,
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.users.insert_one(doc)
    doc.pop("password_hash")
    doc.pop("_id", None)
    return doc

@api.get("/auth/users")
async def list_users(user: dict = Depends(require_role("admin"))):
    users = await db.users.find({}, {"password_hash": 0, "_id": 0}).to_list(1000)
    return users

@api.delete("/auth/users/{uid}")
async def delete_user(uid: str, user: dict = Depends(require_role("admin"))):
    await db.users.delete_one({"id": uid})
    return {"ok": True}

# ----------------- CRUD Helpers -----------------
async def crud_list(coll):
    return [strip_id(d) for d in await db[coll].find().to_list(5000)]

async def crud_create(coll, doc: dict):
    doc = {**doc, "id": doc.get("id") or new_id()}
    await db[coll].insert_one(doc)
    return strip_id(doc)

async def crud_update(coll, id_: str, doc: dict):
    doc.pop("id", None); doc.pop("_id", None)
    await db[coll].update_one({"id": id_}, {"$set": doc})
    return strip_id(await db[coll].find_one({"id": id_}))

async def crud_delete(coll, id_: str):
    await db[coll].delete_one({"id": id_})
    return {"ok": True}

# ----------------- Master Divisi -----------------
@api.get("/divisi")
async def list_divisi(user: dict = Depends(get_current_user)):
    return await crud_list("divisi")

@api.post("/divisi")
async def create_divisi(d: Divisi, user: dict = Depends(require_role("admin"))):
    return await crud_create("divisi", d.model_dump(exclude_none=True))

@api.put("/divisi/{id_}")
async def update_divisi(id_: str, d: Divisi, user: dict = Depends(require_role("admin"))):
    return await crud_update("divisi", id_, d.model_dump(exclude_none=True))

@api.delete("/divisi/{id_}")
async def delete_divisi(id_: str, user: dict = Depends(require_role("admin"))):
    return await crud_delete("divisi", id_)

# ----------------- Master Jabatan -----------------
@api.get("/jabatan")
async def list_jabatan(user: dict = Depends(get_current_user)):
    return await crud_list("jabatan")

@api.post("/jabatan")
async def create_jabatan(d: Jabatan, user: dict = Depends(require_role("admin"))):
    return await crud_create("jabatan", d.model_dump(exclude_none=True))

@api.put("/jabatan/{id_}")
async def update_jabatan(id_: str, d: Jabatan, user: dict = Depends(require_role("admin"))):
    return await crud_update("jabatan", id_, d.model_dump(exclude_none=True))

@api.delete("/jabatan/{id_}")
async def delete_jabatan(id_: str, user: dict = Depends(require_role("admin"))):
    return await crud_delete("jabatan", id_)

# ----------------- Master Karyawan -----------------
@api.get("/karyawan")
async def list_karyawan(user: dict = Depends(get_current_user)):
    return await crud_list("karyawan")

@api.post("/karyawan")
async def create_karyawan(d: Karyawan, user: dict = Depends(require_role("admin"))):
    return await crud_create("karyawan", d.model_dump(exclude_none=True))

@api.put("/karyawan/{id_}")
async def update_karyawan(id_: str, d: Karyawan, user: dict = Depends(require_role("admin"))):
    return await crud_update("karyawan", id_, d.model_dump(exclude_none=True))

@api.delete("/karyawan/{id_}")
async def delete_karyawan(id_: str, user: dict = Depends(require_role("admin"))):
    return await crud_delete("karyawan", id_)

# ----------------- Master KPI -----------------
@api.get("/kpi-master")
async def list_kpi_master(user: dict = Depends(get_current_user)):
    return await crud_list("kpi_master")

@api.post("/kpi-master")
async def create_kpi_master(d: KPIMaster, user: dict = Depends(require_role("admin"))):
    return await crud_create("kpi_master", d.model_dump(exclude_none=True))

@api.put("/kpi-master/{id_}")
async def update_kpi_master(id_: str, d: KPIMaster, user: dict = Depends(require_role("admin"))):
    return await crud_update("kpi_master", id_, d.model_dump(exclude_none=True))

@api.delete("/kpi-master/{id_}")
async def delete_kpi_master(id_: str, user: dict = Depends(require_role("admin"))):
    return await crud_delete("kpi_master", id_)

# ----------------- Target KPI -----------------
@api.get("/kpi-target")
async def list_target(tahun: Optional[int] = None, user: dict = Depends(get_current_user)):
    q = {}
    if tahun: q["tahun"] = tahun
    return [strip_id(d) for d in await db.kpi_target.find(q).to_list(50000)]

@api.post("/kpi-target/bulk")
async def bulk_upsert_target(rows: List[KPITarget], user: dict = Depends(require_role("admin","supervisor"))):
    for r in rows:
        d = r.model_dump(exclude_none=True)
        d["id"] = d.get("id") or new_id()
        await db.kpi_target.update_one(
            {"kode_kpi": d["kode_kpi"], "tahun": d["tahun"], "bulan": d["bulan"]},
            {"$set": {"target": d["target"], "id": d["id"], "kode_kpi": d["kode_kpi"], "tahun": d["tahun"], "bulan": d["bulan"]}},
            upsert=True,
        )
    return {"ok": True, "count": len(rows)}

# ----------------- Input KPI -----------------
@api.get("/kpi-input")
async def list_input(divisi: Optional[str] = None, bulan: Optional[int] = None, tahun: Optional[int] = None,
                     nik: Optional[str] = None, only_approved: bool = False,
                     user: dict = Depends(get_current_user)):
    q = {}
    if bulan: q["bulan"] = bulan
    if tahun: q["tahun"] = tahun
    if nik: q["nik"] = nik
    # karyawan can only see own data
    if user.get("role") == "karyawan":
        if not user.get("nik"):
            return []
        q["nik"] = user["nik"]
    # supervisor scoped to their assigned divisi
    if user.get("role") == "supervisor" and user.get("divisi"):
        divisi = user["divisi"]
    if only_approved:
        q["status"] = "approved"
    inputs = [strip_id(d) for d in await db.kpi_input.find(q).to_list(50000)]
    # enrich with computed fields
    result = []
    karyawan = {k["nik"]: k for k in await crud_list("karyawan")}
    kpi_master = {k["kode"]: k for k in await crud_list("kpi_master")}
    targets = {}
    for t in await db.kpi_target.find({}).to_list(50000):
        targets[(t["kode_kpi"], t["tahun"], t["bulan"])] = t["target"]
    setting = await db.setting.find_one({"id": "default"}) or {}
    on_track_min = setting.get("on_track_min", 85) / 100.0
    for r in inputs:
        km = kpi_master.get(r["kode_kpi"], {})
        kr = karyawan.get(r["nik"], {})
        target = targets.get((r["kode_kpi"], r["tahun"], r["bulan"]), 0)
        real = r.get("realisasi", 0) or 0
        arah = km.get("arah_nilai", "Higher")
        if target == 0 or real == 0:
            ach = 0
        elif arah == "Lower":
            ach = target / real if real else 0
        else:
            ach = real / target if target else 0
        bobot = km.get("bobot", 0)
        nilai = ach * bobot
        if divisi and km.get("divisi") != divisi:
            continue
        r.update({
            "nama": kr.get("nama"),
            "divisi": km.get("divisi"),
            "nama_kpi": km.get("nama"),
            "satuan": km.get("satuan"),
            "arah_nilai": arah,
            "target": target,
            "achievement": round(ach, 4),
            "bobot": bobot,
            "nilai": round(nilai, 4),
            "status_kpi": "On Track" if ach >= on_track_min else "Tertinggal",
            "status": r.get("status", "approved"),
        })
        result.append(r)
    return result

@api.post("/kpi-input/bulk")
async def bulk_upsert_input(rows: List[KPIInput], user: dict = Depends(require_role("admin","supervisor"))):
    default_status = "draft" if user.get("role") == "supervisor" else "approved"
    # Supervisor may only bulk-write for KPIs in their divisi
    if user.get("role") == "supervisor":
        if not user.get("divisi"):
            raise HTTPException(403, "Supervisor belum memiliki divisi. Hubungi Admin untuk assign divisi.")
        kpi_master = {k["kode"]: k for k in await crud_list("kpi_master")}
        for r in rows:
            km = kpi_master.get(r.kode_kpi, {})
            if km.get("divisi") != user["divisi"]:
                raise HTTPException(403, f"Supervisor {user['divisi']} tidak boleh menyimpan KPI divisi {km.get('divisi','?')}")
    for r in rows:
        d = r.model_dump(exclude_none=True)
        d["id"] = d.get("id") or new_id()
        # Server-side authoritative status: supervisor -> draft, admin -> approved
        d["status"] = default_status
        await db.kpi_input.update_one(
            {"nik": d["nik"], "kode_kpi": d["kode_kpi"], "tahun": d["tahun"], "bulan": d["bulan"]},
            {"$set": d}, upsert=True,
        )
    return {"ok": True, "count": len(rows)}

@api.post("/kpi-input/submit")
async def submit_input(body: dict, user: dict = Depends(require_role("supervisor","admin"))):
    """Supervisor submits kpi_input for approval. Body: {tahun, bulan, divisi}"""
    tahun, bulan, divisi = body.get("tahun"), body.get("bulan"), body.get("divisi")
    if user.get("role") == "supervisor" and user.get("divisi"):
        divisi = user["divisi"]
    kpi_master = {k["kode"]: k for k in await crud_list("kpi_master")}
    kodes = [k for k, v in kpi_master.items() if not divisi or v.get("divisi") == divisi]
    res = await db.kpi_input.update_many(
        {"tahun": tahun, "bulan": bulan, "kode_kpi": {"$in": kodes}, "status": {"$in": ["draft", None]}},
        {"$set": {"status": "submitted"}}
    )
    return {"ok": True, "submitted": res.modified_count}

@api.post("/kpi-input/approve")
async def approve_input(body: dict, user: dict = Depends(require_role("admin"))):
    """Admin approves submitted inputs. Body: {tahun, bulan, divisi?, action: approve|reject}"""
    tahun, bulan, divisi = body.get("tahun"), body.get("bulan"), body.get("divisi")
    action = body.get("action", "approve")
    new_status = "approved" if action == "approve" else "rejected"
    kpi_master = {k["kode"]: k for k in await crud_list("kpi_master")}
    kodes = [k for k, v in kpi_master.items() if not divisi or v.get("divisi") == divisi]
    res = await db.kpi_input.update_many(
        {"tahun": tahun, "bulan": bulan, "kode_kpi": {"$in": kodes}, "status": "submitted"},
        {"$set": {"status": new_status}}
    )
    return {"ok": True, "updated": res.modified_count, "status": new_status}

@api.post("/kpi-input/copy-previous")
async def copy_previous_month(body: dict, user: dict = Depends(require_role("admin","supervisor"))):
    """Copy realisasi from previous month. Body: {tahun, bulan, divisi?}"""
    tahun, bulan, divisi = body.get("tahun"), body.get("bulan"), body.get("divisi")
    prev_bulan = bulan - 1
    prev_tahun = tahun
    if prev_bulan < 1:
        prev_bulan = 12
        prev_tahun = tahun - 1
    q = {"tahun": prev_tahun, "bulan": prev_bulan}
    prev = await db.kpi_input.find(q).to_list(50000)
    if divisi:
        kpi_master = {k["kode"]: k for k in await crud_list("kpi_master")}
        prev = [p for p in prev if kpi_master.get(p["kode_kpi"], {}).get("divisi") == divisi]
    count = 0
    default_status = "draft" if user.get("role") == "supervisor" else "approved"
    for p in prev:
        await db.kpi_input.update_one(
            {"nik": p["nik"], "kode_kpi": p["kode_kpi"], "tahun": tahun, "bulan": bulan},
            {"$set": {
                "id": new_id(), "nik": p["nik"], "kode_kpi": p["kode_kpi"],
                "tahun": tahun, "bulan": bulan, "realisasi": p.get("realisasi", 0),
                "status": default_status,
            }},
            upsert=True,
        )
        count += 1
    return {"ok": True, "copied": count, "from": f"{prev_tahun}-{prev_bulan}"}

@api.delete("/kpi-input/{id_}")
async def delete_input(id_: str, user: dict = Depends(require_role("admin","supervisor"))):
    return await crud_delete("kpi_input", id_)

# ----------------- Rekap & Dashboard -----------------
async def get_setting():
    s = await db.setting.find_one({"id": "default"})
    if not s:
        s = {"id": "default", **SettingModel().model_dump()}
        await db.setting.insert_one(s)
    strip_id(s)
    return s

def calc_grade(pct, s):
    if pct >= s["threshold_a"]/100: return "A"
    if pct >= s["threshold_b"]/100: return "B"
    if pct >= s["threshold_c"]/100: return "C"
    if pct >= s["threshold_d"]/100: return "D"
    return "E"

async def compute_rows(bulan: Optional[int], tahun: int, only_approved: bool = True):
    """Return enriched rows for all inputs in bulan/tahun. Rekap uses only approved by default."""
    q = {"tahun": tahun}
    if bulan: q["bulan"] = bulan
    if only_approved:
        q["status"] = {"$in": ["approved", None]}
    inputs = [strip_id(d) for d in await db.kpi_input.find(q).to_list(100000)]
    karyawan = {k["nik"]: k for k in await crud_list("karyawan")}
    kpi_master = {k["kode"]: k for k in await crud_list("kpi_master")}
    targets = {}
    for t in await db.kpi_target.find({"tahun": tahun}).to_list(50000):
        targets[(t["kode_kpi"], t["tahun"], t["bulan"])] = t["target"]
    out = []
    for r in inputs:
        km = kpi_master.get(r["kode_kpi"], {})
        kr = karyawan.get(r["nik"], {})
        target = targets.get((r["kode_kpi"], r["tahun"], r["bulan"]), 0)
        real = r.get("realisasi", 0) or 0
        arah = km.get("arah_nilai", "Higher")
        if target == 0 or real == 0:
            ach = 0
        elif arah == "Lower":
            ach = target / real if real else 0
        else:
            ach = real / target if target else 0
        bobot = km.get("bobot", 0)
        r.update({
            "nama": kr.get("nama"), "divisi": km.get("divisi"),
            "nama_kpi": km.get("nama"), "target": target,
            "achievement": ach, "bobot": bobot, "nilai": ach*bobot,
        })
        out.append(r)
    return out, karyawan, kpi_master

@api.get("/rekap/individu")
async def rekap_individu(tahun: int, bulan: Optional[int] = None, user: dict = Depends(get_current_user)):
    rows, karyawan, _ = await compute_rows(bulan, tahun)
    s = await get_setting()
    # karyawan role: filter own only
    if user.get("role") == "karyawan":
        own_nik = user.get("nik")
        rows = [r for r in rows if r["nik"] == own_nik]
    grouped: Dict[str, Dict[str, Any]] = {}
    for r in rows:
        nik = r["nik"]
        if nik not in grouped:
            kr = karyawan.get(nik, {})
            grouped[nik] = {"nik": nik, "nama": kr.get("nama"), "divisi": kr.get("divisi"),
                            "jabatan": kr.get("jabatan"), "total_nilai": 0, "total_bobot": 0, "items": 0}
        grouped[nik]["total_nilai"] += r["nilai"]
        grouped[nik]["total_bobot"] += r["bobot"]
        grouped[nik]["items"] += 1
    result = []
    for v in grouped.values():
        score = (v["total_nilai"] / v["total_bobot"]) if v["total_bobot"] else 0
        v["score"] = round(score, 4)
        v["grade"] = calc_grade(score, s)
        result.append(v)
    result.sort(key=lambda x: -x["score"])
    return result

@api.get("/rekap/divisi")
async def rekap_divisi(tahun: int, bulan: Optional[int] = None, user: dict = Depends(get_current_user)):
    ind = await rekap_individu(tahun, bulan, user)
    s = await get_setting()
    by_div: Dict[str, Dict[str, Any]] = {}
    for i in ind:
        d = i["divisi"] or "Lain"
        if d not in by_div:
            by_div[d] = {"divisi": d, "sum_score": 0, "count": 0}
        by_div[d]["sum_score"] += i["score"]
        by_div[d]["count"] += 1
    for v in by_div.values():
        v["score"] = round((v["sum_score"] / v["count"]) if v["count"] else 0, 4)
        v["target"] = s["on_track_min"] / 100.0
        v["grade"] = calc_grade(v["score"], s)
    return list(by_div.values())

@api.get("/rekap/perusahaan")
async def rekap_perusahaan(tahun: int, user: dict = Depends(get_current_user)):
    s = await get_setting()
    trend = []
    for m in range(1, 13):
        rows, _, _ = await compute_rows(m, tahun)
        total_nilai = sum(r["nilai"] for r in rows)
        total_bobot = sum(r["bobot"] for r in rows)
        score = (total_nilai/total_bobot) if total_bobot else 0
        trend.append({"bulan": m, "bulan_short": MONTHS_SHORT[m-1], "score": round(score, 4), "target": s["on_track_min"]/100.0})
    return {"trend": trend, "tahun": tahun}

@api.get("/dashboard/overview")
async def dashboard_overview(tahun: int, bulan: int, user: dict = Depends(get_current_user)):
    s = await get_setting()
    ind = await rekap_individu(tahun, bulan, user)
    if user.get("role") == "karyawan":
        # Return simplified personal overview
        my = ind[0] if ind else None
        return {
            "tahun": tahun, "bulan": bulan, "bulan_nama": MONTHS_ID[bulan-1],
            "overall_score": my["score"] if my else 0,
            "target": s["on_track_min"]/100.0,
            "gap": (my["score"] - s["on_track_min"]/100.0) if my else 0,
            "karyawan_dinilai": 1 if my else 0,
            "on_track": 1 if my and my["score"] >= s["on_track_min"]/100.0 else 0,
            "divisi": [], "trend": [],
            "grade_distribusi": {"A":0,"B":0,"C":0,"D":0,"E":0, **({my["grade"]:1} if my else {})},
            "ranking": ind,
            "personal": True,
        }
    divs = await rekap_divisi(tahun, bulan, user)
    per = await rekap_perusahaan(tahun, user)
    overall = 0; total_bobot = 0
    for i in ind:
        overall += i["score"] * i["total_bobot"]
        total_bobot += i["total_bobot"]
    overall_score = (overall/total_bobot) if total_bobot else 0
    grade_dist = {"A":0,"B":0,"C":0,"D":0,"E":0}
    on_track_count = 0
    for i in ind:
        grade_dist[i["grade"]] = grade_dist.get(i["grade"], 0) + 1
        if i["score"] >= s["on_track_min"]/100.0:
            on_track_count += 1
    return {
        "tahun": tahun, "bulan": bulan, "bulan_nama": MONTHS_ID[bulan-1],
        "overall_score": round(overall_score, 4),
        "target": s["on_track_min"]/100.0,
        "gap": round(overall_score - s["on_track_min"]/100.0, 4),
        "karyawan_dinilai": len(ind),
        "on_track": on_track_count,
        "divisi": divs,
        "trend": per["trend"],
        "grade_distribusi": grade_dist,
        "ranking": ind[:10],
    }

@api.get("/reward-punishment")
async def reward_punishment(tahun: int, bulan: Optional[int] = None, user: dict = Depends(get_current_user)):
    if user.get("role") == "karyawan":
        raise HTTPException(403, "Akses ditolak")
    ind = await rekap_individu(tahun, bulan, user)
    for i in ind:
        g = i["grade"]
        if g == "A": i["reward"] = "Bonus + Sertifikat Kinerja Terbaik"
        elif g == "B": i["reward"] = "Bonus Reguler"
        elif g == "C": i["reward"] = "Pembinaan Ringan"
        elif g == "D": i["reward"] = "Peringatan Tertulis"
        else: i["reward"] = "Evaluasi Mendalam / SP"
    return ind

# ----------------- Setting -----------------
@api.get("/setting")
async def get_setting_api(user: dict = Depends(get_current_user)):
    return await get_setting()

@api.put("/setting")
async def update_setting_api(s: SettingModel, user: dict = Depends(require_role("admin"))):
    doc = {"id": "default", **s.model_dump()}
    await db.setting.update_one({"id": "default"}, {"$set": doc}, upsert=True)
    return doc

# ----------------- Payroll -----------------
def _slip_token(nik: str, tahun: int, bulan: int) -> str:
    key = os.environ.get("JWT_SECRET", "").encode()
    msg = f"{nik}|{tahun}|{bulan}".encode()
    sig = hmac.new(key, msg, hashlib.sha256).digest()
    return base64.urlsafe_b64encode(sig)[:16].decode()

def _verify_slip_token(nik: str, tahun: int, bulan: int, token: str) -> bool:
    return hmac.compare_digest(_slip_token(nik, tahun, bulan), token)

async def compute_payroll_row(karyawan: dict, payroll: dict, setting: dict) -> dict:
    gaji_pokok = karyawan.get("gaji_pokok") or 0
    t_trans = karyawan.get("tunjangan_transportasi") or 0
    t_makan = karyawan.get("tunjangan_makan") or 0
    t_kesehatan = karyawan.get("tunjangan_kesehatan") or 0
    total_gaji_kotor = gaji_pokok + t_trans + t_makan + t_kesehatan

    rate_lembur = setting.get("rate_lembur", 12000)
    rate_terlambat = setting.get("rate_terlambat", 12000)
    rate_tdk_masuk = setting.get("rate_tidak_masuk", 36000)

    lembur_jam = payroll.get("lembur_jam") or 0
    terlambat_jam = payroll.get("terlambat_jam") or 0
    tdk_masuk_hari = payroll.get("tidak_masuk_hari") or 0
    pot_pinjaman = payroll.get("potongan_pinjaman") or 0
    pot_lainnya = payroll.get("potongan_lainnya") or 0
    bonus = payroll.get("bonus_lainnya") or 0

    total_lembur = lembur_jam * rate_lembur
    pot_terlambat = terlambat_jam * rate_terlambat
    pot_tdk_masuk = tdk_masuk_hari * rate_tdk_masuk
    total_potongan = pot_terlambat + pot_tdk_masuk + pot_pinjaman + pot_lainnya
    total_lembur_lainnya = total_lembur + bonus
    take_home_pay = total_gaji_kotor + total_lembur_lainnya - total_potongan

    return {
        "gaji_pokok": gaji_pokok,
        "tunjangan_transportasi": t_trans,
        "tunjangan_makan": t_makan,
        "tunjangan_kesehatan": t_kesehatan,
        "total_gaji_kotor": total_gaji_kotor,
        "rate_lembur": rate_lembur,
        "rate_terlambat": rate_terlambat,
        "rate_tidak_masuk": rate_tdk_masuk,
        "total_lembur": total_lembur,
        "bonus_lainnya": bonus,
        "total_lembur_lainnya": total_lembur_lainnya,
        "pot_terlambat": pot_terlambat,
        "pot_tidak_masuk": pot_tdk_masuk,
        "pot_pinjaman": pot_pinjaman,
        "pot_lainnya": pot_lainnya,
        "total_potongan": total_potongan,
        "take_home_pay": take_home_pay,
    }

@api.get("/payroll")
async def list_payroll(tahun: int, bulan: int, user: dict = Depends(get_current_user)):
    if user.get("role") == "karyawan" and not user.get("nik"):
        return []
    kar_list = await crud_list("karyawan")
    if user.get("role") == "karyawan":
        kar_list = [k for k in kar_list if k["nik"] == user["nik"]]
    payrolls = {p["nik"]: strip_id(p) for p in await db.payroll_input.find({"tahun": tahun, "bulan": bulan}).to_list(5000)}
    setting = await get_setting()
    out = []
    for kar in kar_list:
        p = payrolls.get(kar["nik"], {"nik": kar["nik"], "tahun": tahun, "bulan": bulan})
        computed = await compute_payroll_row(kar, p, setting)
        out.append({
            "nik": kar["nik"], "nama": kar["nama"], "jabatan": kar["jabatan"],
            "divisi": kar["divisi"], "sistem_kerja": kar.get("sistem_kerja", "Bulanan"),
            "no_rek": kar.get("no_rek",""), "bank": kar.get("bank","BCA"),
            "periode_gaji": p.get("periode_gaji", ""),
            "catatan": p.get("catatan",""),
            "lembur_jam": p.get("lembur_jam", 0),
            "terlambat_jam": p.get("terlambat_jam", 0),
            "tidak_masuk_hari": p.get("tidak_masuk_hari", 0),
            "potongan_pinjaman": p.get("potongan_pinjaman", 0),
            "potongan_lainnya": p.get("potongan_lainnya", 0),
            **computed,
        })
    out.sort(key=lambda x: x["nama"] or "")
    return out

@api.get("/payroll/{nik}")
async def get_payroll(nik: str, tahun: int, bulan: int, user: dict = Depends(get_current_user)):
    if user.get("role") == "karyawan" and user.get("nik") != nik:
        raise HTTPException(403, "Akses ditolak")
    kar = await db.karyawan.find_one({"nik": nik})
    if not kar: raise HTTPException(404, "Karyawan tidak ditemukan")
    strip_id(kar)
    p = await db.payroll_input.find_one({"nik": nik, "tahun": tahun, "bulan": bulan})
    if p: strip_id(p)
    else: p = {"nik": nik, "tahun": tahun, "bulan": bulan}
    setting = await get_setting()
    computed = await compute_payroll_row(kar, p, setting)
    return {"karyawan": kar, "payroll": p, "hasil": computed}

@api.post("/payroll")
async def upsert_payroll(row: PayrollInput, user: dict = Depends(require_role("admin","supervisor"))):
    d = row.model_dump(exclude_none=True)
    d["id"] = d.get("id") or new_id()
    await db.payroll_input.update_one(
        {"nik": d["nik"], "tahun": d["tahun"], "bulan": d["bulan"]},
        {"$set": d}, upsert=True,
    )
    return {"ok": True}

@api.post("/payroll/bulk")
async def bulk_payroll(rows: List[PayrollInput], user: dict = Depends(require_role("admin","supervisor"))):
    for r in rows:
        d = r.model_dump(exclude_none=True)
        d["id"] = d.get("id") or new_id()
        await db.payroll_input.update_one(
            {"nik": d["nik"], "tahun": d["tahun"], "bulan": d["bulan"]},
            {"$set": d}, upsert=True,
        )
    return {"ok": True, "count": len(rows)}

@api.post("/payroll/copy-previous")
async def copy_payroll_previous(body: dict, user: dict = Depends(require_role("admin","supervisor"))):
    tahun, bulan = body.get("tahun"), body.get("bulan")
    prev_b, prev_y = (bulan-1, tahun) if bulan > 1 else (12, tahun-1)
    prev = await db.payroll_input.find({"tahun": prev_y, "bulan": prev_b}).to_list(5000)
    count = 0
    for p in prev:
        await db.payroll_input.update_one(
            {"nik": p["nik"], "tahun": tahun, "bulan": bulan},
            {"$set": {
                "id": new_id(), "nik": p["nik"], "tahun": tahun, "bulan": bulan,
                "lembur_jam": p.get("lembur_jam", 0),
                "terlambat_jam": p.get("terlambat_jam", 0),
                "tidak_masuk_hari": p.get("tidak_masuk_hari", 0),
                "potongan_pinjaman": p.get("potongan_pinjaman", 0),
                "potongan_lainnya": p.get("potongan_lainnya", 0),
                "bonus_lainnya": p.get("bonus_lainnya", 0),
                "periode_gaji": p.get("periode_gaji", ""),
                "catatan": p.get("catatan", ""),
            }},
            upsert=True,
        )
        count += 1
    return {"ok": True, "copied": count}

@api.get("/payroll/export/rekap")
async def export_payroll_rekap(tahun: int, bulan: int, fmt: str = "excel", user: dict = Depends(get_current_user)):
    if user.get("role") == "karyawan":
        raise HTTPException(403, "Akses ditolak")
    rows = await list_payroll(tahun, bulan, user)
    if fmt == "excel":
        data = [["#","NIK","Nama","Jabatan","Sistem","No Rek","Bank","Take Home Pay"]]
        for i, r in enumerate(rows, 1):
            data.append([i, r["nik"], r["nama"], r["jabatan"], r["sistem_kerja"], r["no_rek"], r["bank"], r["take_home_pay"]])
        total_thp = sum(r["take_home_pay"] for r in rows)
        data.append(["", "", "TOTAL", "", "", "", "", total_thp])
        xlsx = build_workbook({f"Payroll {MONTHS_ID[bulan-1]} {tahun}": data})
        return StreamingResponse(io.BytesIO(xlsx),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="payroll_{tahun}_{bulan:02d}.xlsx"'})
    # PDF rekap
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet(); story = []
    story.append(Paragraph("<b>AP GROUP — Rekap Payroll</b>", styles["Title"]))
    story.append(Paragraph(f"Periode: {MONTHS_ID[bulan-1]} {tahun}", styles["Heading3"]))
    story.append(Spacer(1, 12))
    data = [["#","NIK","Nama","Jabatan","Sistem","Take Home Pay"]]
    for i, r in enumerate(rows, 1):
        data.append([i, r["nik"], r["nama"], r["jabatan"], r["sistem_kerja"], f"Rp {r['take_home_pay']:,.0f}"])
    total_thp = sum(r["take_home_pay"] for r in rows)
    data.append(["", "", "TOTAL", "", "", f"Rp {total_thp:,.0f}"])
    tbl = Table(data, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0), colors.HexColor("#111827")),
        ("TEXTCOLOR",(0,0),(-1,0), colors.white),
        ("FONTSIZE",(0,0),(-1,-1), 9),
        ("GRID",(0,0),(-1,-1),0.3, colors.HexColor("#9CA3AF")),
        ("ALIGN",(-1,0),(-1,-1),"RIGHT"),
        ("BACKGROUND",(0,-1),(-1,-1), colors.HexColor("#F3F4F6")),
        ("FONTNAME",(0,-1),(-1,-1),"Helvetica-Bold"),
    ]))
    story.append(tbl); doc.build(story); buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="payroll_{tahun}_{bulan:02d}.pdf"'})

@api.get("/payroll/slip/public")
async def export_payroll_slip_public(nik: str, tahun: int, bulan: int, token: str, fmt: str = "pdf"):
    """Public tokenized download for WhatsApp sharing (no auth)."""
    if not _verify_slip_token(nik, tahun, bulan, token):
        raise HTTPException(403, "Token tidak valid")
    return await _generate_slip_response(nik, tahun, bulan, fmt)

@api.get("/payroll/slip/{nik}")
async def export_payroll_slip(nik: str, tahun: int, bulan: int, fmt: str = "pdf", user: dict = Depends(get_current_user)):
    if user.get("role") == "karyawan" and user.get("nik") != nik:
        raise HTTPException(403, "Akses ditolak")
    return await _generate_slip_response(nik, tahun, bulan, fmt)

@api.get("/payroll/wa/{nik}")
async def payroll_wa_link(nik: str, tahun: int, bulan: int, user: dict = Depends(require_role("admin","supervisor"))):
    """Return wa.me link + suggested message for sending slip via WhatsApp."""
    kar = await db.karyawan.find_one({"nik": nik})
    if not kar: raise HTTPException(404, "Karyawan tidak ditemukan")
    strip_id(kar)
    p = await db.payroll_input.find_one({"nik": nik, "tahun": tahun, "bulan": bulan}) or {}
    strip_id(p)
    setting = await get_setting()
    h = await compute_payroll_row(kar, p, setting)

    def norm_phone(hp: str) -> str:
        hp = (hp or "").strip().replace(" ","").replace("-","").replace("+","")
        if not hp: return ""
        if hp.startswith("0"): hp = "62" + hp[1:]
        if not hp.startswith("62"): hp = "62" + hp
        return hp

    phone = norm_phone(kar.get("no_hp",""))
    token = _slip_token(nik, tahun, bulan)
    backend = os.environ.get("PUBLIC_BACKEND_URL") or ""
    pdf_url = f"{backend}/api/payroll/slip/public?nik={nik}&tahun={tahun}&bulan={bulan}&token={token}&fmt=pdf"
    fmt_rp = lambda v: f"Rp {int(v or 0):,}".replace(",",".")
    msg = (
        f"*AP GROUP — Slip Gaji {MONTHS_ID[bulan-1]} {tahun}*\n\n"
        f"Halo {kar['nama']},\n"
        f"Berikut ringkasan gaji Anda periode {p.get('periode_gaji') or MONTHS_ID[bulan-1]+' '+str(tahun)}:\n\n"
        f"• Gaji Kotor: {fmt_rp(h['total_gaji_kotor'])}\n"
        f"• Lembur/Bonus: +{fmt_rp(h['total_lembur_lainnya'])}\n"
        f"• Potongan: -{fmt_rp(h['total_potongan'])}\n"
        f"---\n"
        f"*Take Home Pay: {fmt_rp(h['take_home_pay'])}*\n\n"
        f"Slip lengkap (PDF):\n{pdf_url}\n\n"
        f"Terima kasih.\n_HR AP Group_"
    )
    import urllib.parse
    encoded = urllib.parse.quote(msg)
    wa_link = f"https://wa.me/{phone}?text={encoded}" if phone else f"https://wa.me/?text={encoded}"
    return {
        "wa_link": wa_link, "pdf_url": pdf_url, "message": msg,
        "phone": phone, "has_phone": bool(phone),
        "nama": kar["nama"], "take_home_pay": h["take_home_pay"],
    }

async def _generate_slip_response(nik: str, tahun: int, bulan: int, fmt: str):
    kar = await db.karyawan.find_one({"nik": nik})
    if not kar: raise HTTPException(404, "Karyawan tidak ditemukan")
    strip_id(kar)
    p = await db.payroll_input.find_one({"nik": nik, "tahun": tahun, "bulan": bulan}) or {"nik": nik, "tahun": tahun, "bulan": bulan}
    strip_id(p)
    setting = await get_setting()
    h = await compute_payroll_row(kar, p, setting)
    if fmt == "excel":
        rows = [
            ["AP GROUP — SLIP GAJI KARYAWAN"], [""],
            ["A. DATA KARYAWAN"],
            ["Nama Karyawan", ":", kar["nama"]],
            ["Jabatan", ":", kar["jabatan"]],
            ["Periode Gaji", ":", p.get("periode_gaji") or f"{MONTHS_ID[bulan-1]} {tahun}"],
            ["Sistem Kerja", ":", kar.get("sistem_kerja","Bulanan")],
            [""],
            ["B. KOMPONEN GAJI", "Nominal"],
            ["Gaji Pokok", h["gaji_pokok"]],
            ["Tunjangan Transportasi", h["tunjangan_transportasi"]],
            ["Tunjangan Makan", h["tunjangan_makan"]],
            ["Tunjangan Kesehatan", h["tunjangan_kesehatan"]],
            ["Total Gaji Kotor", h["total_gaji_kotor"]], [""],
            ["C. LEMBUR / LAINNYA", "Jumlah", "Satuan", "Harga", "Nominal"],
            ["Lembur", p.get("lembur_jam", 0), "Jam", h["rate_lembur"], h["total_lembur"]],
            ["Bonus Lainnya", "-", "-", "-", h["bonus_lainnya"]],
            ["Total Lembur / Lainnya", "", "", "", h["total_lembur_lainnya"]], [""],
            ["D. POTONGAN GAJI", "Jumlah", "Satuan", "Harga", "Nominal"],
            ["Keterlambatan / Izin", p.get("terlambat_jam", 0), "Jam", h["rate_terlambat"], h["pot_terlambat"]],
            ["Tidak Masuk / Alpha", p.get("tidak_masuk_hari", 0), "Hari", h["rate_tidak_masuk"], h["pot_tidak_masuk"]],
            ["Potongan Pinjaman", "-", "-", "-", h["pot_pinjaman"]],
            ["Potongan Lainnya", "-", "-", "-", h["pot_lainnya"]],
            ["Total Potongan", "", "", "", h["total_potongan"]], [""],
            ["E. REKAP GAJI"],
            ["Total Gaji Kotor", h["total_gaji_kotor"]],
            ["Total Lembur / Lainnya", h["total_lembur_lainnya"]],
            ["Total Potongan", -h["total_potongan"]],
            ["Gaji Diterima (Take Home Pay)", h["take_home_pay"]], [""],
            ["F. CATATAN", p.get("catatan","")],
        ]
        xlsx = build_workbook({"Slip Gaji": rows})
        return StreamingResponse(io.BytesIO(xlsx),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="slip_{nik}_{tahun}_{bulan:02d}.xlsx"'})
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet(); story = []
    story.append(Paragraph("<b>AP GROUP</b>", styles["Title"]))
    story.append(Paragraph("SLIP GAJI KARYAWAN", styles["Heading2"]))
    story.append(Spacer(1, 12))
    data_karyawan = [
        ["Nama Karyawan", kar["nama"]],
        ["Jabatan", kar["jabatan"]],
        ["Periode Gaji", p.get("periode_gaji") or f"{MONTHS_ID[bulan-1]} {tahun}"],
        ["Sistem Kerja", kar.get("sistem_kerja","Bulanan")],
        ["No Rekening", f"{kar.get('no_rek','')} ({kar.get('bank','')})"],
    ]
    t = Table(data_karyawan, colWidths=[150, 350])
    t.setStyle(TableStyle([("FONTSIZE",(0,0),(-1,-1),9),("GRID",(0,0),(-1,-1),0.3, colors.HexColor("#9CA3AF"))]))
    story.append(Paragraph("<b>A. DATA KARYAWAN</b>", styles["Heading4"])); story.append(t); story.append(Spacer(1, 12))
    fmt_rp = lambda v: f"Rp {v:,.0f}"
    comp = [["Komponen","Nominal"],
            ["Gaji Pokok", fmt_rp(h["gaji_pokok"])],
            ["Tunjangan Transportasi", fmt_rp(h["tunjangan_transportasi"])],
            ["Tunjangan Makan", fmt_rp(h["tunjangan_makan"])],
            ["Tunjangan Kesehatan", fmt_rp(h["tunjangan_kesehatan"])],
            ["Total Gaji Kotor", fmt_rp(h["total_gaji_kotor"])]]
    tt = Table(comp, colWidths=[350, 150])
    tt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0), colors.HexColor("#111827")),("TEXTCOLOR",(0,0),(-1,0), colors.white),
        ("FONTSIZE",(0,0),(-1,-1),9),("GRID",(0,0),(-1,-1),0.3, colors.HexColor("#9CA3AF")),
        ("ALIGN",(-1,0),(-1,-1),"RIGHT"),("FONTNAME",(0,-1),(-1,-1),"Helvetica-Bold")]))
    story.append(Paragraph("<b>B. KOMPONEN GAJI</b>", styles["Heading4"])); story.append(tt); story.append(Spacer(1, 12))
    lembur = [["Komponen","Jumlah","Satuan","Harga","Nominal"],
              ["Lembur", p.get("lembur_jam", 0), "Jam", fmt_rp(h["rate_lembur"]), fmt_rp(h["total_lembur"])],
              ["Bonus Lainnya", "-", "-", "-", fmt_rp(h["bonus_lainnya"])],
              ["Total Lembur / Lainnya", "", "", "", fmt_rp(h["total_lembur_lainnya"])]]
    tl = Table(lembur, colWidths=[200, 60, 60, 90, 90])
    tl.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0), colors.HexColor("#111827")),("TEXTCOLOR",(0,0),(-1,0), colors.white),
        ("FONTSIZE",(0,0),(-1,-1),9),("GRID",(0,0),(-1,-1),0.3, colors.HexColor("#9CA3AF")),("FONTNAME",(0,-1),(-1,-1),"Helvetica-Bold")]))
    story.append(Paragraph("<b>C. LEMBUR / LAINNYA</b>", styles["Heading4"])); story.append(tl); story.append(Spacer(1, 12))
    pot = [["Jenis","Jumlah","Satuan","Harga","Jumlah Potongan"],
           ["Keterlambatan / Izin", p.get("terlambat_jam", 0), "Jam", fmt_rp(h["rate_terlambat"]), fmt_rp(h["pot_terlambat"])],
           ["Tidak Masuk / Alpha", p.get("tidak_masuk_hari", 0), "Hari", fmt_rp(h["rate_tidak_masuk"]), fmt_rp(h["pot_tidak_masuk"])],
           ["Potongan Pinjaman", "-", "-", "-", fmt_rp(h["pot_pinjaman"])],
           ["Potongan Lainnya", "-", "-", "-", fmt_rp(h["pot_lainnya"])],
           ["Total Potongan", "", "", "", fmt_rp(h["total_potongan"])]]
    tp = Table(pot, colWidths=[200, 60, 60, 90, 90])
    tp.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0), colors.HexColor("#DC2626")),("TEXTCOLOR",(0,0),(-1,0), colors.white),
        ("FONTSIZE",(0,0),(-1,-1),9),("GRID",(0,0),(-1,-1),0.3, colors.HexColor("#9CA3AF")),("FONTNAME",(0,-1),(-1,-1),"Helvetica-Bold")]))
    story.append(Paragraph("<b>D. POTONGAN GAJI</b>", styles["Heading4"])); story.append(tp); story.append(Spacer(1, 12))
    rekap = [["Keterangan","Nominal"],
             ["Total Gaji Kotor", fmt_rp(h["total_gaji_kotor"])],
             ["Total Lembur / Lainnya", fmt_rp(h["total_lembur_lainnya"])],
             ["Total Potongan", f"- {fmt_rp(h['total_potongan'])}"],
             ["GAJI DITERIMA (TAKE HOME PAY)", fmt_rp(h["take_home_pay"])]]
    tr = Table(rekap, colWidths=[350, 150])
    tr.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0), colors.HexColor("#111827")),("TEXTCOLOR",(0,0),(-1,0), colors.white),
        ("FONTSIZE",(0,0),(-1,-1),10),("GRID",(0,0),(-1,-1),0.3, colors.HexColor("#9CA3AF")),
        ("ALIGN",(-1,0),(-1,-1),"RIGHT"),
        ("BACKGROUND",(0,-1),(-1,-1), colors.HexColor("#059669")),
        ("TEXTCOLOR",(0,-1),(-1,-1), colors.white),
        ("FONTNAME",(0,-1),(-1,-1),"Helvetica-Bold")]))
    story.append(Paragraph("<b>E. REKAP GAJI</b>", styles["Heading4"])); story.append(tr); story.append(Spacer(1, 12))
    if p.get("catatan"):
        story.append(Paragraph("<b>F. CATATAN</b>", styles["Heading4"]))
        story.append(Paragraph(p["catatan"], styles["BodyText"]))
    doc.build(story); buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="slip_{nik}_{tahun}_{bulan:02d}.pdf"'})

# ----------------- YoY -----------------
@api.get("/rekap/yoy")
async def yoy_comparison(tahun: int, years: int = 3, user: dict = Depends(get_current_user)):
    """Return multi-year monthly trend comparison"""
    result = []
    for y in range(tahun - years + 1, tahun + 1):
        r = await rekap_perusahaan(y, user)
        result.append({"tahun": y, "trend": r["trend"]})
    return {"years": result, "focus_year": tahun}

# ----------------- Export -----------------
def build_workbook(sheets: Dict[str, List[List[Any]]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

@api.get("/export/excel")
async def export_excel(tahun: int, bulan: Optional[int] = None, jenis: str = "individu", user: dict = Depends(get_current_user)):
    if user.get("role") == "karyawan" and jenis != "individu":
        raise HTTPException(403, "Akses ditolak")
    if jenis == "individu":
        data = await rekap_individu(tahun, bulan, user)
        rows = [["NIK","Nama","Divisi","Jabatan","Total Nilai","Score","Grade"]]
        for d in data:
            rows.append([d["nik"], d["nama"], d["divisi"], d["jabatan"], round(d["total_nilai"],2), round(d["score"]*100,2), d["grade"]])
    elif jenis == "divisi":
        data = await rekap_divisi(tahun, bulan, user)
        rows = [["Divisi","Score (%)","Target (%)","Grade"]]
        for d in data:
            rows.append([d["divisi"], round(d["score"]*100,2), round(d["target"]*100,2), d["grade"]])
    elif jenis == "perusahaan":
        per = await rekap_perusahaan(tahun, user)
        rows = [["Bulan","Aktual (%)","Target (%)"]]
        for t in per["trend"]:
            rows.append([MONTHS_ID[t["bulan"]-1], round(t["score"]*100,2), round(t["target"]*100,2)])
    elif jenis == "reward":
        data = await reward_punishment(tahun, bulan, user)
        rows = [["NIK","Nama","Divisi","Score (%)","Grade","Reward/Punishment"]]
        for d in data:
            rows.append([d["nik"], d["nama"], d["divisi"], round(d["score"]*100,2), d["grade"], d["reward"]])
    else:
        raise HTTPException(400, "Jenis laporan tidak dikenali")
    xlsx = build_workbook({f"Rekap {jenis}": rows})
    fn = f"rekap_{jenis}_{tahun}_{bulan or 'all'}.xlsx"
    return StreamingResponse(io.BytesIO(xlsx),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fn}"'})

@api.get("/export/pdf")
async def export_pdf(tahun: int, bulan: Optional[int] = None, jenis: str = "individu", user: dict = Depends(get_current_user)):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    story = []
    title = f"Laporan Rekap {jenis.capitalize()} — {MONTHS_ID[bulan-1] if bulan else 'Tahunan'} {tahun}"
    story.append(Paragraph(f"<b>AP GROUP — KPI Management</b>", styles["Title"]))
    story.append(Paragraph(title, styles["Heading3"]))
    story.append(Spacer(1, 12))
    if user.get("role") == "karyawan" and jenis != "individu":
        raise HTTPException(403, "Akses ditolak")
    if jenis == "individu":
        data = await rekap_individu(tahun, bulan, user)
        rows = [["NIK","Nama","Divisi","Jabatan","Total Nilai","Score (%)","Grade"]]
        for d in data:
            rows.append([d["nik"], d["nama"], d["divisi"], d["jabatan"], f"{d['total_nilai']:.2f}", f"{d['score']*100:.2f}", d["grade"]])
    elif jenis == "divisi":
        data = await rekap_divisi(tahun, bulan, user)
        rows = [["Divisi","Score (%)","Target (%)","Grade"]]
        for d in data:
            rows.append([d["divisi"], f"{d['score']*100:.2f}", f"{d['target']*100:.2f}", d["grade"]])
    elif jenis == "perusahaan":
        per = await rekap_perusahaan(tahun, user)
        rows = [["Bulan","Aktual (%)","Target (%)"]]
        for t in per["trend"]:
            rows.append([MONTHS_ID[t["bulan"]-1], f"{t['score']*100:.2f}", f"{t['target']*100:.2f}"])
    elif jenis == "reward":
        data = await reward_punishment(tahun, bulan, user)
        rows = [["NIK","Nama","Divisi","Score (%)","Grade","Reward / Punishment"]]
        for d in data:
            rows.append([d["nik"], d["nama"], d["divisi"], f"{d['score']*100:.2f}", d["grade"], d["reward"]])
    else:
        raise HTTPException(400, "Jenis laporan tidak dikenali")
    table = Table(rows, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0), colors.HexColor("#111827")),
        ("TEXTCOLOR",(0,0),(-1,0), colors.white),
        ("FONTSIZE",(0,0),(-1,-1), 9),
        ("GRID",(0,0),(-1,-1),0.3, colors.HexColor("#9CA3AF")),
        ("ALIGN",(0,0),(-1,-1),"LEFT"),
        ("PADDING",(0,0),(-1,-1),4),
    ]))
    story.append(table)
    doc.build(story)
    buf.seek(0)
    fn = f"rekap_{jenis}_{tahun}_{bulan or 'all'}.pdf"
    return StreamingResponse(buf, media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{fn}"'})

# ----------------- Excel Import -----------------
def _clean_drawings(xlsx_bytes: bytes) -> bytes:
    """Strip drawings/charts/media/rels for safe openpyxl load."""
    import zipfile, tempfile, shutil
    tmp = tempfile.mkdtemp()
    src_path = os.path.join(tmp, "in.xlsx")
    with open(src_path, "wb") as f: f.write(xlsx_bytes)
    ext_dir = os.path.join(tmp, "ext")
    os.makedirs(ext_dir, exist_ok=True)
    with zipfile.ZipFile(src_path, "r") as z:
        z.extractall(ext_dir)
    # remove drawing/chart/media dirs entirely
    for sub in ("xl/drawings","xl/charts","xl/media","xl/embeddings"):
        p = os.path.join(ext_dir, sub)
        if os.path.exists(p): shutil.rmtree(p)
    # clean sheet rels (drawing, chart, image, vmlDrawing) - full tag regex
    rels_dir = os.path.join(ext_dir, "xl/worksheets/_rels")
    if os.path.exists(rels_dir):
        for f in os.listdir(rels_dir):
            fp = os.path.join(rels_dir, f)
            with open(fp) as fh: c = fh.read()
            c = re.sub(r'<Relationship\b[^>]*Type="[^"]*(drawing|chart|image|vmlDrawing)"[^>]*/>', '', c)
            with open(fp, "w") as fh: fh.write(c)
    # clean drawing tag inside sheet xml
    ws_dir = os.path.join(ext_dir, "xl/worksheets")
    for f in os.listdir(ws_dir):
        if f.endswith(".xml"):
            fp = os.path.join(ws_dir, f)
            with open(fp) as fh: c = fh.read()
            c = re.sub(r'<drawing[^/]*/>','', c)
            c = re.sub(r'<legacyDrawing[^/]*/>','', c)
            c = re.sub(r'<picture[^/]*/>','', c)
            with open(fp, "w") as fh: fh.write(c)
    # clean workbook rels
    wb_rels = os.path.join(ext_dir, "xl/_rels/workbook.xml.rels")
    if os.path.exists(wb_rels):
        with open(wb_rels) as fh: c = fh.read()
        c = re.sub(r'<Relationship\b[^>]*Type="[^"]*(chart|drawing|image|vbaProject)"[^>]*/>', '', c)
        with open(wb_rels, "w") as fh: fh.write(c)
    # clean chartsheet dir if any
    cs_dir = os.path.join(ext_dir, "xl/chartsheets")
    if os.path.exists(cs_dir): shutil.rmtree(cs_dir)
    # remove content type overrides
    ct = os.path.join(ext_dir, "[Content_Types].xml")
    with open(ct) as fh: c = fh.read()
    c = re.sub(r'<Override[^>]*(drawing|chart|image|vbaProject|vmlDrawing)[^"]*"[^/]*/>', '', c)
    c = re.sub(r'<Default[^>]*(png|jpg|jpeg|gif|bmp|emf|wmf|bin)"[^/]*/>', '', c)
    with open(ct, "w") as fh: fh.write(c)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for root, dirs, files in os.walk(ext_dir):
            for f in files:
                full = os.path.join(root, f)
                rel = os.path.relpath(full, ext_dir)
                z.write(full, rel)
    shutil.rmtree(tmp)
    return out.getvalue()

async def import_from_xlsx_bytes(xlsx_bytes: bytes) -> dict:
    clean = _clean_drawings(xlsx_bytes)
    wb = openpyxl.load_workbook(io.BytesIO(clean), data_only=True)
    stats = {}
    # Master Divisi
    if "02_Master Divisi" in wb.sheetnames:
        ws = wb["02_Master Divisi"]
        await db.divisi.delete_many({})
        rows = list(ws.iter_rows(min_row=5, values_only=True))
        docs = []
        for r in rows:
            if r and r[0] and r[1]:
                docs.append({"id": new_id(), "kode": str(r[0]), "nama": str(r[1])})
        if docs: await db.divisi.insert_many(docs)
        stats["divisi"] = len(docs)
    # Master Jabatan
    if "03_Master Jabatan" in wb.sheetnames:
        ws = wb["03_Master Jabatan"]
        await db.jabatan.delete_many({})
        docs = []
        for r in ws.iter_rows(min_row=5, values_only=True):
            if r and r[0] and r[1]:
                docs.append({"id": new_id(), "kode": str(r[0]), "nama": str(r[1]), "divisi": str(r[2] or "")})
        if docs: await db.jabatan.insert_many(docs)
        stats["jabatan"] = len(docs)
    # Master Karyawan
    if "04_Master Karyawan" in wb.sheetnames:
        ws = wb["04_Master Karyawan"]
        await db.karyawan.delete_many({})
        docs = []
        for r in ws.iter_rows(min_row=5, values_only=True):
            if r and r[0] and r[1]:
                docs.append({"id": new_id(), "nik": str(r[0]), "nama": str(r[1]),
                             "divisi": str(r[2] or ""), "jabatan": str(r[3] or ""),
                             "atasan": str(r[4] or "Owner"), "status": str(r[5] or "Aktif"),
                             "gaji_pokok": 0, "tunjangan_transportasi": 0,
                             "tunjangan_makan": 0, "tunjangan_kesehatan": 0,
                             "sistem_kerja": "Bulanan", "no_rek": "", "bank": "BCA"})
        if docs: await db.karyawan.insert_many(docs)
        stats["karyawan"] = len(docs)
    # Master KPI
    if "05_Master KPI" in wb.sheetnames:
        ws = wb["05_Master KPI"]
        await db.kpi_master.delete_many({})
        docs = []
        for r in ws.iter_rows(min_row=5, values_only=True):
            if r and r[0] and r[1] and r[2]:
                try:
                    bobot = float(r[4] or 0)
                except Exception: bobot = 0
                docs.append({"id": new_id(), "divisi": str(r[0]), "kode": str(r[1]),
                             "nama": str(r[2]), "satuan": str(r[3] or ""), "bobot": bobot,
                             "arah_nilai": str(r[5] or "Higher"), "jenis_target": str(r[6] or "Rasio")})
        if docs: await db.kpi_master.insert_many(docs)
        stats["kpi_master"] = len(docs)
    # Target KPI - matrix Jan..Dec
    if "06_Target KPI" in wb.sheetnames:
        ws = wb["06_Target KPI"]
        # first read tahun from row 3 col B
        tahun_cell = ws.cell(3, 2).value
        try: tahun = int(tahun_cell)
        except Exception: tahun = datetime.now().year
        await db.kpi_target.delete_many({"tahun": tahun})
        docs = []
        for r in ws.iter_rows(min_row=5, values_only=True):
            if not r or not r[1]: continue
            kode = str(r[1])
            for m in range(12):
                try:
                    tgt = float(r[4+m]) if r[4+m] is not None else 0
                except Exception: tgt = 0
                docs.append({"id": new_id(), "kode_kpi": kode, "tahun": tahun, "bulan": m+1, "target": tgt})
        if docs: await db.kpi_target.insert_many(docs)
        stats["kpi_target"] = len(docs)
    # KPI Input from sheets 07A..07F
    input_sheets = [s for s in wb.sheetnames if s.startswith("07")]
    bulan_map = {n:i+1 for i,n in enumerate(MONTHS_ID)}
    await db.kpi_input.delete_many({})
    total_input = 0
    for sn in input_sheets:
        ws = wb[sn]
        docs = []
        for r in ws.iter_rows(min_row=5, values_only=True):
            if not r or not r[0]: continue
            bulan = bulan_map.get(str(r[0]).strip())
            if not bulan: continue
            try: tahun = int(r[1])
            except Exception: continue
            nik = str(r[3] or "")
            kode = str(r[5] or "")
            try: real = float(r[8]) if r[8] is not None else 0
            except Exception: real = 0
            if not nik or not kode: continue
            docs.append({"id": new_id(), "bulan": bulan, "tahun": tahun, "nik": nik,
                         "kode_kpi": kode, "realisasi": real,
                         "cut_off": str(r[2]) if r[2] else None})
        if docs: await db.kpi_input.insert_many(docs)
        total_input += len(docs)
    stats["kpi_input"] = total_input
    return stats

@api.post("/import/excel")
async def import_excel(file: UploadFile = File(...), user: dict = Depends(require_role("admin"))):
    content = await file.read()
    try:
        stats = await import_from_xlsx_bytes(content)
    except Exception as e:
        logger.exception("Import failed")
        raise HTTPException(400, f"Gagal import: {e}")
    return {"ok": True, "stats": stats}

@api.post("/import/payroll")
async def import_payroll_from_excel(file: UploadFile = File(...), user: dict = Depends(require_role("admin"))):
    """Import employees + payroll fields from AP Group payroll Excel file (uses 'Preview' sheet)."""
    content = await file.read()
    try:
        clean = _clean_drawings(content)
        wb = openpyxl.load_workbook(io.BytesIO(clean), data_only=True)
    except Exception as e:
        raise HTTPException(400, f"Gagal baca file: {e}")

    # Find preview sheet
    preview_sheet = None
    for sn in wb.sheetnames:
        if "preview" in sn.lower() and "copy" not in sn.lower():
            preview_sheet = sn; break
    if not preview_sheet:
        for sn in wb.sheetnames:
            if "preview" in sn.lower():
                preview_sheet = sn; break
    if not preview_sheet:
        raise HTTPException(400, "Sheet Preview payroll tidak ditemukan")
    ws = wb[preview_sheet]

    # Preview rows: col B = Nama, C = Jabatan, D = Before25, E = THP Non Rev, G = No Rek, H = Bank
    imported = 0; updated_thp = 0
    existing = {k["nama"].lower().strip(): k for k in await crud_list("karyawan") if k.get("nama")}
    for r in ws.iter_rows(min_row=3, values_only=True):
        if not r or not r[1]: continue
        nama = str(r[1]).strip()
        if nama.upper().startswith("TOTAL"): continue
        jabatan = str(r[2] or "").strip()
        try: thp = float(r[4] or r[3] or 0)
        except Exception: thp = 0
        no_rek = str(r[6] or "").strip() if len(r) > 6 else ""
        bank = str(r[7] or "BCA").strip() if len(r) > 7 else "BCA"
        # match by name (case-insensitive prefix)
        matched = None
        for existing_nama, kar in existing.items():
            if existing_nama[:15] == nama.lower()[:15] or nama.lower().startswith(existing_nama[:15]):
                matched = kar; break
        if matched:
            await db.karyawan.update_one({"nik": matched["nik"]},
                {"$set": {"jabatan": jabatan or matched["jabatan"], "no_rek": no_rek, "bank": bank,
                          "gaji_pokok": thp}})
            updated_thp += 1
        else:
            # create new karyawan (nik = generated from nama)
            nik = "AUTO" + str(uuid.uuid4())[:6].upper()
            div = "Lain"
            # try to infer divisi from jabatan text
            for dname in ["Marketing","Produksi","Admin","Design","HR","Finance"]:
                if dname.lower() in jabatan.lower(): div = dname; break
            await db.karyawan.insert_one({
                "id": new_id(), "nik": nik, "nama": nama, "divisi": div,
                "jabatan": jabatan, "atasan": "Owner", "status": "Aktif",
                "gaji_pokok": thp, "tunjangan_transportasi": 0, "tunjangan_makan": 0,
                "tunjangan_kesehatan": 0, "sistem_kerja": "Bulanan",
                "no_rek": no_rek, "bank": bank,
            })
            imported += 1
    return {"ok": True, "karyawan_baru": imported, "karyawan_updated": updated_thp}


    """Re-import the built-in AP Group template."""
    path = os.environ.get("KPI_TEMPLATE_PATH", "/app/AP_GROUP_KPI_template.xlsm")
    if not os.path.exists(path):
        raise HTTPException(404, f"Template tidak ditemukan: {path}")
    with open(path, "rb") as f:
        stats = await import_from_xlsx_bytes(f.read())
    return {"ok": True, "stats": stats}

# ----------------- Startup -----------------
@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.users.create_index("id")
    for c in ("divisi","jabatan","karyawan","kpi_master","kpi_target","kpi_input","setting"):
        await db[c].create_index("id")
    # seed admin
    admin_email = os.environ["ADMIN_EMAIL"].lower()
    admin_pw = os.environ["ADMIN_PASSWORD"]
    existing = await db.users.find_one({"email": admin_email})
    if not existing:
        await db.users.insert_one({
            "id": new_id(), "email": admin_email, "name": os.environ.get("ADMIN_NAME","Admin"),
            "role": "admin", "password_hash": hash_pw(admin_pw),
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        logger.info(f"Admin seeded: {admin_email}")
    else:
        if not verify_pw(admin_pw, existing["password_hash"]):
            await db.users.update_one({"email": admin_email}, {"$set": {"password_hash": hash_pw(admin_pw)}})
    # seed setting
    await get_setting()
    # auto import template if empty
    if (await db.divisi.count_documents({})) == 0:
        path = os.environ.get("KPI_TEMPLATE_PATH", "/app/AP_GROUP_KPI_template.xlsm")
        if os.path.exists(path):
            try:
                with open(path, "rb") as f:
                    stats = await import_from_xlsx_bytes(f.read())
                logger.info(f"Auto imported template: {stats}")
            except Exception:
                logger.exception("Auto import failed")
    # auto import payroll master (gaji_pokok, no_rek) if not yet done
    if (await db.karyawan.count_documents({"gaji_pokok": {"$gt": 0}})) == 0:
        payroll_path = "/app/AP_GROUP_payroll_template.xlsx"
        if os.path.exists(payroll_path):
            try:
                # simulate call
                clean = _clean_drawings(open(payroll_path,"rb").read())
                wb = openpyxl.load_workbook(io.BytesIO(clean), data_only=True)
                preview_sheet = None
                for sn in wb.sheetnames:
                    if "preview" in sn.lower() and "copy" not in sn.lower():
                        preview_sheet = sn; break
                if preview_sheet:
                    ws = wb[preview_sheet]
                    existing = {k["nama"].lower().strip(): k for k in await crud_list("karyawan") if k.get("nama")}
                    imported = updated_thp = 0
                    for r in ws.iter_rows(min_row=3, values_only=True):
                        if not r or not r[1]: continue
                        nama = str(r[1]).strip()
                        if nama.upper().startswith("TOTAL"): continue
                        jabatan = str(r[2] or "").strip()
                        try: thp = float(r[4] or r[3] or 0)
                        except Exception: thp = 0
                        no_rek = str(r[6] or "").strip() if len(r) > 6 else ""
                        bank = str(r[7] or "BCA").strip() if len(r) > 7 else "BCA"
                        matched = None
                        for en, kar in existing.items():
                            if en[:15] == nama.lower()[:15] or nama.lower().startswith(en[:15]):
                                matched = kar; break
                        if matched:
                            await db.karyawan.update_one({"nik": matched["nik"]},
                                {"$set": {"no_rek": no_rek, "bank": bank, "gaji_pokok": thp}})
                            updated_thp += 1
                        else:
                            nik = "AUTO" + str(uuid.uuid4())[:6].upper()
                            div = "Lain"
                            for dname in ["Marketing","Produksi","Admin","Design","HR","Finance"]:
                                if dname.lower() in jabatan.lower(): div = dname; break
                            await db.karyawan.insert_one({
                                "id": new_id(), "nik": nik, "nama": nama, "divisi": div,
                                "jabatan": jabatan, "atasan": "Owner", "status": "Aktif",
                                "gaji_pokok": thp, "tunjangan_transportasi": 0, "tunjangan_makan": 0,
                                "tunjangan_kesehatan": 0, "sistem_kerja": "Bulanan",
                                "no_rek": no_rek, "bank": bank,
                            })
                            imported += 1
                    logger.info(f"Auto imported payroll master: baru={imported}, updated={updated_thp}")
            except Exception:
                logger.exception("Payroll auto import failed")
    # seed demo users if not exist
    for uinfo in [
        {"email":"supervisor@apgroup.com","name":"Supervisor Demo","role":"supervisor","password":"Supervisor123!"},
        {"email":"karyawan@apgroup.com","name":"Karyawan Demo","role":"karyawan","password":"Karyawan123!","nik":"EMP001"},
    ]:
        if not await db.users.find_one({"email": uinfo["email"]}):
            await db.users.insert_one({
                "id": new_id(), "email": uinfo["email"], "name": uinfo["name"],
                "role": uinfo["role"], "password_hash": hash_pw(uinfo["password"]),
                "nik": uinfo.get("nik"),
                "created_at": datetime.now(timezone.utc).isoformat(),
            })

@app.on_event("shutdown")
async def shutdown():
    client.close()

@api.get("/")
async def root():
    return {"app": "AP GROUP KPI Management API", "status": "ok"}

frontend_url = os.environ.get('FRONTEND_URL')
_cors = os.environ.get('CORS_ORIGINS','*').split(',')
if frontend_url and frontend_url not in _cors: _cors.append(frontend_url)
app.include_router(api)
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=_cors,
    allow_methods=["*"],
    allow_headers=["*"],
)
